"""Build one Word document per viewport from the Playwright screenshot folder.

Each screenshot is written with a title and description above the image.
Output files:

    screenshots/APEX-screenshots-desktop.docx
    screenshots/APEX-screenshots-mobile.docx
    screenshots/APEX-screenshots-tablet.docx
"""

from __future__ import annotations

import struct
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SCREENSHOT_DIR = ROOT / "screenshots"
OUTPUT_DIR = SCREENSHOT_DIR
VIEWPORTS = ("desktop", "mobile", "tablet")
MAX_IMAGE_WIDTH_IN = 6.4
MAX_IMAGE_HEIGHT_IN = 8.2

# Display order. Files not listed here are appended alphabetically.
ORDER = [
    "students-dashboard",
    "students-dashboard--mobile-nav",
    "students-dashboard--submission-tracker-details",
    "students-dashboard--submission-tracker-progress",
    "students-dashboard--review-notice-denied",
    "students-dashboard--review-notice-returned",
    "students-dashboard--sign-out",
    "students-submissions",
    "students-submissions--guide",
    "students-saaf-step-1-classification",
    "students-saaf-step-2-proponents",
    "students-saaf-step-3-details",
    "students-saaf-step-4-alignment",
    "students-saaf--confirm-clear",
    "students-saaf--confirm-submit",
    "students-saaf--success",
    "students-reservations",
    "students-reservations--confirm-clear",
    "students-reservations--confirm-submit",
    "students-about",
    "signatories-dashboard",
    "signatories-dashboard--filter",
    "signatories-dashboard--activity-detail",
    "signatories-dashboard--approve-confirm",
    "signatories-dashboard--return-proposal",
    "signatories-dashboard--reject-proposal",
    "signatories-dashboard--mobile-nav",
    "signatories-dashboard--sign-out",
    "signatories-about",
    "admin-dashboard",
    "admin-dashboard--submission-detail",
    "admin-dashboard--new-announcement",
    "admin-dashboard--edit-announcement",
    "admin-dashboard--delete-announcement",
    "admin-organizations",
    "admin-organizations--edit",
    "admin-signatories",
    "admin-signatories--edit",
    "admin-dashboard--mobile-nav",
    "admin-dashboard--sign-out",
    "admin-about",
]

DESCRIPTIONS: dict[str, tuple[str, str]] = {
    "students-dashboard": (
        "Student organization dashboard",
        "Home screen for an organization officer after sign-in. The table lists "
        "filed SAAF proposals with document ID, event title, classification, "
        "current signatory, and status. Announcements from OSA sit below the "
        "table. Timeline reminders on the right surface denied papers, returned "
        "papers, and upcoming deadlines. Create Project/Event starts a new filing.",
    ),
    "students-dashboard--mobile-nav": (
        "Student mobile navigation",
        "Top sheet used on phone and tablet widths instead of the persistent "
        "sidebar. It shows the APEX brand, Dashboard and New proposal links, "
        "Sign Out, and the signed-in officer name, role, and About APEX link.",
    ),
    "students-dashboard--submission-tracker-details": (
        "Submission tracker — Document Details",
        "Opened by tapping a row in Project Status & Submissions. The Details "
        "tab shows the official packet: event name, classification, venue, "
        "schedule, expected attendees, budget, description, and related fields "
        "from the filed SAAF.",
    ),
    "students-dashboard--submission-tracker-progress": (
        "Submission tracker — Milestone & Progress",
        "Second tab of the tracker modal. It shows signatory routing for this "
        "paper (adviser, dean, CDM when a venue is reserved, OSAAR) so officers "
        "can see who currently holds the file and which desks have already acted.",
    ),
    "students-dashboard--review-notice-denied": (
        "Denied review notice",
        "Opened from a Denied reminder on the dashboard. It names the activity, "
        "the signatory who denied it, and the rejection comment. A denied paper "
        "cannot be edited; the officer must file a new proposal if they still "
        "want to run the activity.",
    ),
    "students-dashboard--review-notice-returned": (
        "Returned review notice",
        "Opened from a Returned reminder. It shows the signatory's revision "
        "comment. The officer can reopen the SAAF from the tracker (when the "
        "status is returned) and resubmit after making the requested changes.",
    ),
    "students-dashboard--sign-out": (
        "Student sign-out confirmation",
        "Confirms that the officer wants to leave APEX. Cancel keeps the "
        "session. Sign out ends the Cognito session and returns the browser to "
        "the hosted login.",
    ),
    "students-submissions": (
        "Create activity proposal — event setup",
        "First step of a new filing. The officer enters the official event name "
        "and chooses whether school facilities will be reserved. Continue opens "
        "the Student Activity Application Form. The Guide button explains the "
        "three-step workflow.",
    ),
    "students-submissions--guide": (
        "Submission guide",
        "How-it-works modal for filing a proposal: event setup, complete the "
        "SAAF (and reservation form when facilities are needed), then submit "
        "and track multi-level signatory review.",
    ),
    "students-saaf-step-1-classification": (
        "SAAF step 1 — Activity classification",
        "First page of the Student Activity Application Form. The officer "
        "marks the activity as co-curricular or extra-curricular and enters "
        "the total number of class or organization members. Continue moves to "
        "proponents.",
    ),
    "students-saaf-step-2-proponents": (
        "SAAF step 2 — Proponents",
        "People page of the SAAF. Each proponent needs name, student number, "
        "program and year, department, position, organization or section, "
        "contact number, email, and Facebook URL. Additional proponents can "
        "be added before continuing.",
    ),
    "students-saaf-step-3-details": (
        "SAAF step 3 — Activity details",
        "Event description page: title and nature, description (minimum 100 "
        "characters), objectives, venue, start and end dates, times, expected "
        "participants, individual contribution, and proposed budget. The event "
        "date must be at least 11 days after the date of submission.",
    ),
    "students-saaf-step-4-alignment": (
        "SAAF step 4 — Institutional alignment and budget",
        "Final SAAF page. The officer selects at least one Mapúa mission "
        "statement and explains core values, program educational objectives "
        "(optional), and UN Sustainable Development Goals. A line-item budget "
        "and grand total sit below. From here the officer can clear the form, "
        "save a PDF, submit, or go to the reservation form when facilities "
        "were requested.",
    ),
    "students-saaf--confirm-clear": (
        "SAAF confirm clear",
        "Warns that clearing the Student Activity Application Form discards "
        "every field the officer has entered. Cancel keeps the draft. Clear "
        "resets the form to empty defaults.",
    ),
    "students-saaf--confirm-submit": (
        "SAAF confirm submit",
        "Last check before the SAAF is sent into signatory routing. The action "
        "cannot be undone. Cancel returns to the form. Proceed posts the packet.",
    ),
    "students-saaf--success": (
        "SAAF submitted successfully",
        "Shown after the backend accepts the activity application. Close "
        "returns the officer to a fresh form so they can start another filing "
        "or go back to the dashboard to track this one.",
    ),
    "students-reservations": (
        "Reservation of facilities",
        "Optional second form used when the officer chose to reserve school "
        "facilities. It covers equipment (chairs, boards, tables, and others), "
        "general facilities such as the North or South Circle, function rooms, "
        "and audiovisual equipment with dates, times, and locations. Submit "
        "files the combined SAAF and reservation packet.",
    ),
    "students-reservations--confirm-clear": (
        "Reservation confirm clear",
        "Warns that clearing the facility reservation form removes equipment "
        "checks and all facility, room, and AV rows. Cancel keeps the draft.",
    ),
    "students-reservations--confirm-submit": (
        "Reservation confirm submit",
        "Confirms that the officer is filing the SAAF together with the "
        "facility reservation. Proceed sends the combined packet into routing.",
    ),
    "students-about": (
        "About APEX — student desk",
        "Product and team page reached from the sidebar. It explains APEX "
        "(Administrative Portal For Events Exchange), the SAAF-to-signatory "
        "pipeline, and the student developers behind the portal. The student "
        "sidebar remains available for Dashboard and New proposal.",
    ),
    "signatories-dashboard": (
        "Signatory review dashboard",
        "Queue for a signed-in signatory (adviser, dean, CDM reviewer, or "
        "OSAAR). Each row is a paper currently on that desk, with organization, "
        "activity name, submitted date, activity type, and a decision cue "
        "(Review, Return, or Reject). Opening a row shows the full proposal.",
    ),
    "signatories-dashboard--filter": (
        "Signatory department and organization filter",
        "Dropdown that narrows the review queue by department and, once a "
        "department is chosen, by organization. Clearing the filter restores "
        "the full queue for this signatory.",
    ),
    "signatories-dashboard--activity-detail": (
        "Activity proposal detail",
        "Review modal for one paper. The header shows organization, submitted "
        "date, title, and representative. The body lists proponents, schedule, "
        "venue, budget, participant count, description, and objectives. Footer "
        "actions are Close, Return Proposal, Reject Proposal, and Approve Proposal.",
    ),
    "signatories-dashboard--approve-confirm": (
        "Approve proposal confirmation",
        "Nested confirm dialog from Approve Proposal. Proceed records this "
        "signatory's endorsement and advances the paper to the next desk, or "
        "marks it fully approved when OSAAR endorses.",
    ),
    "signatories-dashboard--return-proposal": (
        "Return proposal for revision",
        "Nested dialog for sending the paper back to the student organization. "
        "A return comment is required. The submission stays on this signatory's "
        "desk with Returned status until the officer resubmits.",
    ),
    "signatories-dashboard--reject-proposal": (
        "Reject proposal",
        "Nested dialog for a final denial. A rejection comment is required. "
        "Denied papers leave the queue and cannot be edited by the student.",
    ),
    "signatories-dashboard--mobile-nav": (
        "Signatory mobile navigation",
        "Top navigation sheet on phone and tablet for the signatory desk. It "
        "exposes Dashboard, Sign Out, and About APEX, plus the signed-in "
        "reviewer's name and role.",
    ),
    "signatories-dashboard--sign-out": (
        "Signatory sign-out confirmation",
        "Confirms that the reviewer wants to leave APEX. Cancel keeps the "
        "session so they can continue endorsing or returning papers.",
    ),
    "signatories-about": (
        "About APEX — signatory desk",
        "Same About content as other roles, framed in the signatory shell so "
        "a reviewer can read the product story without leaving the review desk.",
    ),
    "admin-dashboard": (
        "Admin panel — Office of Student Affairs",
        "OSA overview of every organization's submissions, with status and "
        "activity-type filters. Clicking a row opens the full SAAF record. "
        "The announcements table below is the campus bulletin: officers see "
        "these notices on their dashboards. New announcement posts a notice.",
    ),
    "admin-dashboard--submission-detail": (
        "Admin submission detail",
        "Read-only SAAF snapshot for administrators: document ID, status, "
        "current signatory, venue, date, budget, submitted time, and the "
        "notification / routing stepper for that paper.",
    ),
    "admin-dashboard--new-announcement": (
        "New announcement",
        "Compose dialog for posting a bulletin notice. Content is required "
        "(up to 5,000 characters). Saving publishes the notice to student "
        "organization dashboards.",
    ),
    "admin-dashboard--edit-announcement": (
        "Edit announcement",
        "Opened from an existing bulletin row. OSA can replace the notice "
        "text or start a delete. Save changes updates the bulletin in place.",
    ),
    "admin-dashboard--delete-announcement": (
        "Delete announcement confirmation",
        "Nested alert that the notice will be removed from the bulletin and "
        "cannot be undone. Keep cancels. Delete removes it from every "
        "organization dashboard.",
    ),
    "admin-organizations": (
        "Organizations directory",
        "Admin page for registering student organizations. The left column "
        "adds one organization (name, dean, adviser) or imports a CSV. The "
        "right column lists registered organizations. Admin, CDM, and OSAAR "
        "are shared campus accounts, not chosen per organization.",
    ),
    "admin-organizations--edit": (
        "Edit organization",
        "Updates an existing organization's display name and the dean and "
        "adviser desks assigned to it. The organization ID is assigned by "
        "the API and cannot be changed.",
    ),
    "admin-signatories": (
        "Signatories directory",
        "Admin page for registering people who sit on review desks. Add a "
        "name and role (dean, adviser, admin, CDM, OSAAR). Department applies "
        "only to deans. CSV import is available for batch registration. The "
        "table on the right lists everyone already in the directory.",
    ),
    "admin-signatories--edit": (
        "Edit signatory",
        "Updates a registered signatory's name, role, and (for deans) "
        "department. Use this when a desk holder changes or a name needs "
        "correction after import.",
    ),
    "admin-dashboard--mobile-nav": (
        "Admin mobile navigation",
        "Top navigation sheet on phone and tablet for OSA. It lists Dashboard, "
        "Organizations, Signatories, Sign Out, and About APEX.",
    ),
    "admin-dashboard--sign-out": (
        "Admin sign-out confirmation",
        "Confirms that the OSA administrator wants to leave APEX. Cancel "
        "keeps the session so directory and bulletin work can continue.",
    ),
    "admin-about": (
        "About APEX — admin desk",
        "About page from the admin shell. Same product and team story, with "
        "the OSA sidebar still available for Dashboard, Organizations, and "
        "Signatories.",
    ),
}

ROLE_SECTIONS = [
    ("students-", "Student organization officer"),
    ("signatories-", "Signatory reviewer"),
    ("admin-", "OSA administrator"),
]


def png_size(path: Path) -> tuple[int, int]:
    with path.open("rb") as handle:
        signature = handle.read(8)
        if signature != b"\x89PNG\r\n\x1a\n":
            raise ValueError(f"{path} is not a PNG file")
        length, chunk_type = struct.unpack(">I4s", handle.read(8))
        if chunk_type != b"IHDR" or length < 8:
            raise ValueError(f"{path} is missing a PNG IHDR chunk")
        width, height = struct.unpack(">II", handle.read(8))
        return width, height


def image_size_inches(path: Path) -> tuple[float, float]:
    width_px, height_px = png_size(path)
    ratio = width_px / height_px
    width = MAX_IMAGE_WIDTH_IN
    height = width / ratio
    if height > MAX_IMAGE_HEIGHT_IN:
        height = MAX_IMAGE_HEIGHT_IN
        width = height * ratio
    return width, height


def title_for(stem: str) -> str:
    if stem in DESCRIPTIONS:
        return DESCRIPTIONS[stem][0]
    return stem.replace("--", " — ").replace("-", " ").capitalize()


def description_for(stem: str, viewport: str) -> str:
    if stem in DESCRIPTIONS:
        text = DESCRIPTIONS[stem][1]
    else:
        text = (
            f"Captured screen `{stem}` from the APEX {viewport} pass. "
            "Open the live app to inspect the corresponding page or modal."
        )
    return text


def set_run_font(run, name: str = "Calibri") -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def add_horizontal_line(paragraph) -> None:
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_before = Pt(4)
    paragraph_format.space_after = Pt(10)
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "8B0000")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def sort_stems(stems: list[str]) -> list[str]:
    rank = {stem: index for index, stem in enumerate(ORDER)}
    return sorted(stems, key=lambda stem: (rank.get(stem, len(ORDER)), stem))


def role_label(stem: str) -> str | None:
    for prefix, label in ROLE_SECTIONS:
        if stem.startswith(prefix):
            return label
    return None


def build_document(viewport: str, files: list[Path]) -> Document:
    document = Document()

    for section in document.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    heading = document.add_heading(
        f"APEX screenshots — {viewport.capitalize()}",
        level=0,
    )
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

    intro = document.add_paragraph()
    run = intro.add_run(
        f"This file collects every captured {viewport} screen from APEX "
        f"(Administrative Portal For Events Exchange). Each entry is one "
        f"screenshot type: a page, a modal, or a nested modal step. The "
        f"description sits above the image."
    )
    set_run_font(run)
    run.font.size = Pt(11)
    intro.paragraph_format.space_after = Pt(16)

    current_role: str | None = None
    for index, path in enumerate(files, start=1):
        stem = path.stem
        role = role_label(stem)
        if role and role != current_role:
            current_role = role
            section_heading = document.add_heading(role, level=1)
            for run in section_heading.runs:
                run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

        title = document.add_heading(f"{index}. {title_for(stem)}", level=2)
        title.paragraph_format.space_before = Pt(10)
        title.paragraph_format.space_after = Pt(4)

        body = document.add_paragraph()
        run = body.add_run(description_for(stem, viewport))
        set_run_font(run)
        run.font.size = Pt(11)
        body.paragraph_format.space_after = Pt(8)

        width, height = image_size_inches(path)
        document.add_picture(str(path), width=Inches(width), height=Inches(height))
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cap_run = caption.add_run(f"")
        set_run_font(cap_run)
        cap_run.font.size = Pt(9)
        cap_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        cap_run.italic = True
        add_horizontal_line(caption)

    return document


def collect_files(viewport: str) -> list[Path]:
    folder = SCREENSHOT_DIR / viewport
    if not folder.is_dir():
        raise FileNotFoundError(f"Missing screenshot folder: {folder}")
    files = [path for path in folder.glob("*.png") if not path.name.startswith("_")]
    if not files:
        raise FileNotFoundError(f"No PNG screenshots in {folder}")
    by_stem = {path.stem: path for path in files}
    return [by_stem[stem] for stem in sort_stems(list(by_stem))]


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for viewport in VIEWPORTS:
        files = collect_files(viewport)
        document = build_document(viewport, files)
        output = OUTPUT_DIR / f"APEX-screenshots-{viewport}.docx"
        document.save(output)
        written.append(output)
        print(f"Wrote {output} ({len(files)} screenshots)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
